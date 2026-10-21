"""Fetch full-text sources for the gold set and add coverage analysis.

Usage:
  python python/scripts/fetch_gold_sources.py \
      --gold docs/remediation/gold_set/gold_set.jsonl \
      --out docs/remediation/gold_set/sources \
      [--tagged-xlsx docs/remediation/india_team_package/04_AIRI_Review_Feedback_gold_tagged.xlsx] \
      [--delay 3] [--limit N] [--skip-existing]

For arXiv records it downloads https://arxiv.org/pdf/<id> (3 s between requests,
per arXiv's robots guidance), extracts text with pdfminer, and stores
<out>/<risk_id>.txt. For other URLs it fetches the HTML and extracts the article
text with trafilatura when available, otherwise a plain tag strip.

It then writes <out>/../gold_set_sources_stats.json and, when --tagged-xlsx is
given, appends per-record columns to the 'All Feedback' sheet:
  Full Text Words, Record Description Words, Description / Paper Ratio, Has Limitations Section,
  Has Ethics/Impact Section, Has Validation/Evaluation Section,
  Mitigation Terms in Full Text, Fetch Status.

Requires outbound access to arxiv.org (and the other source hosts).
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import time
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

UA = "AIRiskIntellect-gold-set-fetch/1.0 (contact: engineering@accelerateai.io)"
ARXIV_RE = re.compile(r"arxiv\.org/(?:abs|pdf)/([0-9]{4}\.[0-9]{4,5}(?:v\d+)?)")

SECTION_PATTERNS = {
    "limitations": re.compile(r"^\s*(\d+(\.\d+)*\s+)?(limitations?|limitations and future work|discussion and limitations)\b", re.I | re.M),
    "ethics_impact": re.compile(r"^\s*(\d+(\.\d+)*\s+)?(ethic(s|al)( statement| considerations)?|broader impacts?|impact statement|societal impact|responsible ai)\b", re.I | re.M),
    "validation": re.compile(r"^\s*(\d+(\.\d+)*\s+)?(experiments?|evaluation|results|validation|empirical (study|results|evaluation))\b", re.I | re.M),
}
MITIGATION_TERMS = re.compile(r"\b(mitigat\w*|safeguard\w*|guardrail\w*|de-?identif\w*|anonymi[sz]\w*|human[- ]in[- ]the[- ]loop|oversight|validation|audit\w*|red[- ]team\w*|filter\w*|access control|consent|IRB|ethics (approval|board))\b", re.I)


def fetch(url: str, timeout: int = 60) -> tuple[int, bytes, str]:
    req = Request(url, headers={"User-Agent": UA, "Accept": "application/pdf,text/html;q=0.9,*/*;q=0.8"})
    try:
        with urlopen(req, timeout=timeout) as resp:  # noqa: S310 (URLs come from our own gold set)
            return resp.status, resp.read(), resp.headers.get("Content-Type", "")
    except HTTPError as e:
        return e.code, b"", ""
    except URLError as e:
        return 0, b"", str(e.reason)


def pdf_text(data: bytes) -> str:
    from pdfminer.high_level import extract_text
    return extract_text(io.BytesIO(data)) or ""


def html_text(data: bytes) -> str:
    html = data.decode("utf-8", errors="replace")
    try:
        import trafilatura  # type: ignore
        out = trafilatura.extract(html, include_comments=False, include_tables=False)
        if out:
            return out
    except Exception:
        pass
    html = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S | re.I)
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html)).strip()


def analyse(text: str, abstract_words: int) -> dict:
    words = len(text.split())
    return {
        "full_text_words": words,
        "abstract_words": abstract_words,
        "coverage_ratio": round(abstract_words / words, 3) if words else None,
        "has_limitations_section": bool(SECTION_PATTERNS["limitations"].search(text)),
        "has_ethics_impact_section": bool(SECTION_PATTERNS["ethics_impact"].search(text)),
        "has_validation_section": bool(SECTION_PATTERNS["validation"].search(text)),
        "mitigation_terms_in_full_text": len(MITIGATION_TERMS.findall(text)),
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--gold", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--tagged-xlsx")
    ap.add_argument("--delay", type=float, default=3.0)
    ap.add_argument("--limit", type=int)
    ap.add_argument("--skip-existing", action="store_true")
    args = ap.parse_args()

    out = Path(args.out); out.mkdir(parents=True, exist_ok=True)
    gold = [json.loads(l) for l in open(args.gold, encoding="utf-8")]
    if args.limit:
        gold = gold[: args.limit]

    results: dict[str, dict] = {}
    for i, g in enumerate(gold, 1):
        rid, url = g["risk_id"], g["article_url"]
        txt_path = out / f"{rid}.txt"
        rec = {"risk_id": rid, "url": url, "fetch_status": "", "source_kind": ""}
        m = ARXIV_RE.search(url)
        if args.skip_existing and txt_path.exists():
            text = txt_path.read_text(encoding="utf-8")
            rec["fetch_status"] = "cached"; rec["source_kind"] = "arxiv_pdf" if m else "html"
        else:
            if m:
                target = f"https://arxiv.org/pdf/{m.group(1)}"; rec["source_kind"] = "arxiv_pdf"
            else:
                target = url; rec["source_kind"] = "html"
            status, data, ctype = fetch(target)
            if status != 200 or not data:
                rec["fetch_status"] = f"failed:{status or ctype}"; results[rid] = rec
                print(f"[{i}/{len(gold)}] {rid} FAILED {target} ({rec['fetch_status']})", file=sys.stderr)
                time.sleep(args.delay); continue
            try:
                text = pdf_text(data) if (m or "pdf" in ctype.lower()) else html_text(data)
            except Exception as e:  # noqa: BLE001
                rec["fetch_status"] = f"parse_error:{type(e).__name__}"; results[rid] = rec
                time.sleep(args.delay); continue
            txt_path.write_text(text, encoding="utf-8")
            rec["fetch_status"] = "ok"
            print(f"[{i}/{len(gold)}] {rid} ok {len(text.split())} words", file=sys.stderr)
            time.sleep(args.delay)
        abstract_words = len(str(g.get("record", {}).get("Description") or "").split())
        rec.update(analyse(text, abstract_words))
        results[rid] = rec

    ok = [r for r in results.values() if r["fetch_status"] in ("ok", "cached")]
    stats = {
        "records": len(results), "fetched_ok": len(ok),
        "failed": sorted(r["risk_id"] for r in results.values() if r["fetch_status"] not in ("ok", "cached")),
        "median_full_text_words": sorted(r["full_text_words"] for r in ok)[len(ok) // 2] if ok else None,
        "with_limitations_section": sum(r["has_limitations_section"] for r in ok),
        "with_ethics_impact_section": sum(r["has_ethics_impact_section"] for r in ok),
        "with_validation_section": sum(r["has_validation_section"] for r in ok),
        "mean_mitigation_terms": round(sum(r["mitigation_terms_in_full_text"] for r in ok) / len(ok), 1) if ok else None,
        "per_record": results,
    }
    stats_path = out.parent / "gold_set_sources_stats.json"
    stats_path.write_text(json.dumps(stats, indent=2))
    print(json.dumps({k: v for k, v in stats.items() if k != "per_record"}, indent=2))

    if args.tagged_xlsx:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = load_workbook(args.tagged_xlsx)
        ws = wb["All Feedback"]
        hdr = [c.value for c in ws[1]]
        cols = ["Full Text Words", "Record Description Words", "Description / Paper Ratio", "Has Limitations Section", "Has Ethics/Impact Section",
                "Has Validation/Evaluation Section", "Mitigation Terms in Full Text", "Fetch Status"]
        keys = ["full_text_words", "abstract_words", "coverage_ratio", "has_limitations_section", "has_ethics_impact_section",
                "has_validation_section", "mitigation_terms_in_full_text", "fetch_status"]
        start = ws.max_column + 1 if "Full Text Words" not in hdr else hdr.index("Full Text Words") + 1
        for i, name in enumerate(cols):
            c = ws.cell(row=1, column=start + i, value=name)
            c.fill = PatternFill("solid", fgColor="FF1F3A5F"); c.font = Font(bold=True, color="FFFFFFFF"); c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(start + i)].width = 14
        rid_col = hdr.index("Risk ID") + 1
        for r in range(2, ws.max_row + 1):
            rec = results.get(ws.cell(row=r, column=rid_col).value)
            if not rec:
                continue
            for i, k in enumerate(keys):
                v = rec.get(k)
                ws.cell(row=r, column=start + i, value=("Yes" if v is True else "No" if v is False else v)).font = Font(name="Calibri", size=11)
        ws.auto_filter.ref = f"A1:{get_column_letter(start + len(cols) - 1)}{ws.max_row}"
        wb.save(args.tagged_xlsx)
        print("updated", args.tagged_xlsx)
    return 0 if len(ok) == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())
